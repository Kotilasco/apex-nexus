from datetime import date
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Apex_Nexus_System_Functions_E2E_Guide.docx"

FRONTEND_URL = "http://localhost:3000"
GATEWAY_URL = "http://localhost:8200/api"
POSTGRES_CMD = "docker exec apex-postgres psql -U apex_admin -d apex_nexus"
ADMIN_USER = "admin"
ADMIN_PASSWORD = "Admin@2024!"


def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)

    for level in range(1, 5):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0x17, 0x3D, 0x63)


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr[idx].text = str(header)
        for paragraph in hdr[idx].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            for paragraph in cells[idx].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()


def bullet(doc: Document, text: str):
    doc.add_paragraph(text, style="List Bullet")


def code_block(doc: Document, lines):
    for line in lines:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    doc.add_paragraph()


def add_title_page(doc: Document):
    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("APEX NEXUS")
    run.bold = True
    run.font.size = Pt(32)
    run.font.color.rgb = RGBColor(0x17, 0x3D, 0x63)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("System Functions And End-To-End Test Guide")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x4B, 0x67, 0x85)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Generated {date.today().strftime('%d %B %Y')}")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x75, 0x75, 0x75)

    for _ in range(3):
        doc.add_paragraph()

    purpose = doc.add_paragraph()
    purpose.alignment = WD_ALIGN_PARAGRAPH.CENTER
    purpose.add_run(
        "This document describes the business functions available in Apex Nexus and the practical steps to validate each module end to end in a live environment."
    )
    doc.add_page_break()


def add_intro(doc: Document):
    doc.add_heading("1. Scope", level=1)
    doc.add_paragraph(
        "Apex Nexus is a microservice-based enterprise content management platform. This guide is written for product owners, QA engineers, implementation teams, and operations staff who need one Word document that explains what the system does and how to prove each function works."
    )
    doc.add_paragraph(
        "The coverage below is organised around the actual frontend navigation and API routes in this repository. Each module section contains a short purpose statement, the main user-visible functions, and a practical end-to-end validation path."
    )

    doc.add_heading("2. Environment Reference", level=1)
    add_table(doc, ["Item", "Value"], [
        ["Web UI", FRONTEND_URL],
        ["API Gateway", GATEWAY_URL],
        ["PostgreSQL", "localhost:5460"],
        ["Elasticsearch", "http://localhost:9200"],
        ["MinIO Console", "http://localhost:9031"],
        ["MailHog", "http://localhost:8625"],
        ["Admin user", ADMIN_USER],
        ["Admin password", ADMIN_PASSWORD],
    ])

    doc.add_heading("3. Common Test Prerequisites", level=1)
    for item in [
        "Run docker compose up -d and wait until the gateway, auth, document, workflow, search, retention, audit, and notification services are healthy.",
        "Use the admin account for full-path tests. Register secondary users when a scenario needs peer review, approval, signature, or forwarded work.",
        "Keep one sample PDF, one DOCX, one image, and one scanned image-only PDF available for upload tests.",
        "For API tests, obtain a JWT first and send it as Authorization: Bearer <token>.",
    ]:
        bullet(doc, item)

    doc.add_heading("4. Token Acquisition", level=1)
    code_block(doc, [
        f"TOKEN=$(curl -s -X POST {GATEWAY_URL}/auth/login \\",
        "  -H 'Content-Type: application/json' \\",
        f"  -d '{{\"username\":\"{ADMIN_USER}\",\"password\":\"{ADMIN_PASSWORD}\"}}' | jq -r '.data.accessToken // .data.token // .token')",
        "echo $TOKEN",
    ])

    doc.add_heading("5. Core Lifecycle Smoke Flow", level=1)
    for item in [
        "Login, upload a document, confirm it appears in Documents, and inspect version 1 metadata.",
        "Search the document by title in Search and by natural-language phrasing in Semantic Search.",
        "Start a workflow, approve or reject it from a second account, and verify the state transition in Workflow History.",
        "Apply a retention setting or legal hold, then confirm the action appears in Audit Log.",
        "Open the same document in two browser sessions to validate presence and notification behaviour.",
    ]:
        bullet(doc, item)


def add_module_catalog(doc: Document):
    doc.add_heading("6. Module Catalog", level=1)
    rows = [
        ["Dashboard", "/dashboard", "Operational summary, quick actions, trend charts"],
        ["Documents", "/documents", "Upload, preview, versioning, notes, legal hold, checkout, sharing"],
        ["Intake", "/intake", "Capture, classification, ingestion queue"],
        ["Agents", "/agents", "Intent detection, autonomous suggestions, multi-agent controls"],
        ["Cases", "/cases", "Adaptive case folders, tasks, participants, evidence attachment"],
        ["Sustainability", "/sustainability", "Green index, deduplication, archive savings, carbon metrics"],
        ["Federated Search", "/federated", "Unified search across remote systems"],
        ["Projects", "/projects", "Project scoping, membership, AI and retention defaults"],
        ["Workflow", "/workflow", "Definitions, instances, approvals, escalations, forwarding"],
        ["Predictions", "/predictions", "AI-driven forecasting and anomaly views"],
        ["Vendor Portals", "/vendor-portals", "External upload portals with compliance controls"],
        ["Search", "/search", "Keyword, advanced, semantic, OCR, GDPR scan, ask/synthesis"],
        ["Apps / Marketplace", "/apps and /marketplace", "Plugin discovery and activation"],
        ["Knowledge Graph", "/knowledge-graph", "Relationship visualisation across entities"],
        ["Notifications", "/notifications", "Unread queue, read state, event delivery"],
        ["Retention", "/retention", "Policies, scans, disposition queue"],
        ["Compliance", "/compliance", "POTRAZ / privacy / DSR / consent / breach workflows"],
        ["Workflow Designer", "/workflow-designer", "Design workflow definitions and rules"],
        ["Industry", "/industry", "Industry template browsing"],
        ["SAP Integration", "/sap", "P2P, M2C, ArchiveLink, SAP mock services"],
        ["Email Ingestion", "/email-ingestion", "Mailbox configs, rules, polling"],
        ["Trust Center", "/trust-center", "AI governance, transparency, trust controls"],
        ["Analytics", "/analytics", "Operational metrics and workflow health"],
        ["Audit", "/audit", "Immutable audit review and filtering"],
        ["Admin", "/admin/users, /admin/settings", "User admin, role admin, system settings"],
    ]
    add_table(doc, ["Module", "UI Route", "Primary Function"], rows)


def add_module_section(doc: Document, title: str, route: str, apis, functions, e2e, evidence=None):
    doc.add_heading(title, level=2)
    doc.add_paragraph(f"Primary route: {route}")
    if apis:
        doc.add_paragraph("Key APIs:")
        for api in apis:
            bullet(doc, api)
    doc.add_paragraph("Functions:")
    for item in functions:
        bullet(doc, item)
    doc.add_paragraph("End-to-end validation:")
    for step in e2e:
        bullet(doc, step)
    if evidence:
        doc.add_paragraph("Useful verification points:")
        for item in evidence:
            bullet(doc, item)


def build_modules(doc: Document):
    doc.add_heading("7. Detailed Functions And E2E Tests", level=1)

    modules = [
        {
            "title": "7.1 Authentication, Users, And Roles",
            "route": "/login, /register, /admin/users, /admin/settings",
            "apis": ["POST /auth/login", "POST /auth/register", "GET /auth/me", "GET/PUT /auth/users/**", "GET /auth/roles"],
            "functions": [
                "Login, logout, and JWT session enforcement.",
                "User registration and profile bootstrap.",
                "Role assignment, user activation/deactivation, and RBAC enforcement.",
                "Admin maintenance of users and system settings.",
            ],
            "e2e": [
                "Login as admin and verify redirect to Dashboard plus full sidebar visibility.",
                "Register a new user, login as that user, and confirm reduced navigation consistent with AUTHOR permissions.",
                "Promote that user from Admin > Users, refresh their session, and verify newly granted modules appear.",
                "Disable the user and confirm subsequent login is blocked.",
            ],
            "evidence": ["JWT stored in localStorage as apex_token.", "User row updated in users and user_roles tables."]
        },
        {
            "title": "7.2 Dashboard",
            "route": "/dashboard",
            "apis": ["Dashboard data is aggregated from document, workflow, retention, audit, and notification services."],
            "functions": [
                "High-level operational KPIs.",
                "Weekly and monthly trend charts.",
                "Quick-action launch points into major workflows.",
            ],
            "e2e": [
                "Open Dashboard after login and verify KPI cards render without console errors.",
                "Click each quick action and confirm navigation to the target module.",
                "Upload a document and complete a workflow, then refresh the dashboard to confirm counts move.",
            ],
        },
        {
            "title": "7.3 Documents, Folders, And Sharing",
            "route": "/documents",
            "apis": ["GET/POST /documents", "GET/POST /folders/**", "POST /documents/{id}/checkout", "POST /documents/{id}/checkin", "POST /documents/{id}/legal-hold", "POST /auth/share-links"],
            "functions": [
                "Document upload, metadata editing, deletion, and folder organisation.",
                "Version history, compare, download per version, and redacted content retrieval.",
                "Checkout/checkin, cancellation, locks, WOPI token generation, and preview URLs.",
                "Notes, links between documents, legal holds, retention updates, and external share links.",
            ],
            "e2e": [
                "Upload a PDF and DOCX from Documents or Documents/Create and confirm both appear in the list with version 1.",
                "Create a folder, move or upload a document into it, and verify listing by folder works.",
                "Checkout a document, edit/check in a new version, and compare versions.",
                "Add a note, create a document-to-document link, and generate a share link with expiry.",
                "Apply legal hold, then inspect Audit to confirm the hold action was logged.",
            ],
            "evidence": ["documents, document_versions, document_notes, share_links, and audit_log rows are created or updated."]
        },
        {
            "title": "7.4 Search, OCR, Classification, GDPR Scan, And Ask",
            "route": "/search",
            "apis": ["GET/POST /search", "GET /search/semantic", "POST /search/extract", "POST /search/classify", "POST /search/gdpr-scan", "POST /search/ask", "POST /search/synthesize", "POST /search/intent"],
            "functions": [
                "Keyword and advanced metadata search.",
                "Semantic/vector search and natural-language question answering.",
                "OCR/content extraction, document classification, and GDPR/PII scanning.",
                "Language detection for English, Shona, and Ndebele within classification responses.",
            ],
            "e2e": [
                "Search for a recently uploaded document by title and tag.",
                "Run Semantic Search using a natural-language query and verify ranked results differ from literal search.",
                "Upload a scanned PDF to OCR and confirm extract text is returned.",
                "Classify a Shona or Ndebele sample and confirm the response includes language=sn or language=nd.",
                "Run GDPR scan on a file with email addresses or phone numbers and confirm entities are detected.",
            ],
            "evidence": ["Elasticsearch indexes updated; classify-text/classify responses include label, category, confidence, and language."]
        },
        {
            "title": "7.5 Intake",
            "route": "/intake",
            "apis": ["/intake endpoints behind document service", "Search-service classification APIs are called by intake processing."],
            "functions": [
                "Landing zone for captured files before or during processing.",
                "Automatic classification, tagging, and project suggestion.",
                "Hand-off into document creation plus downstream agent triggers.",
            ],
            "e2e": [
                "Submit a file through Intake and confirm it moves from queued to classified/processed.",
                "Verify the created document inherits label, tags, and language.",
                "Open Agents afterwards and confirm agentic suggestions were emitted for the new intake item where applicable.",
            ],
        },
        {
            "title": "7.6 Agentic AI And Multi-Agent Team",
            "route": "/agents",
            "apis": ["GET /agentic/suggestions", "GET /agentic/suggestions/history", "GET /agentic/summary", "GET /agentic/team", "POST /agentic/detect/{documentId}", "POST /agentic/team/{auditor|archivist|bridge}/run"],
            "functions": [
                "Intent detection for invoice auto-match, budget variance, contract renewal, meter anomaly, and compliance review.",
                "Auditor agent for PII/compliance findings, Archivist agent for cold-storage candidates, and Bridge agent for SAP draft preparation.",
                "User approval or dismissal of proposed actions.",
            ],
            "e2e": [
                "Open Agents and verify Pending, Accepted, and Dismissed counts load.",
                "Run detect on a document or upload an invoice/contract that should trigger a suggestion.",
                "Accept one suggestion and dismiss another, then verify they move from Pending to History.",
                "Use the Multi-Agent Team buttons to run Auditor, Archivist, and Bridge and confirm team counters update.",
            ],
            "evidence": ["agentic_suggestions contains status and agent_name; audit_log records AI-related actions."]
        },
        {
            "title": "7.7 Adaptive Case Management",
            "route": "/cases",
            "apis": ["GET/POST /cases", "GET /cases/summary", "GET /cases/{id}", "POST /cases/{id}/status", "POST /cases/{id}/tasks", "POST /cases/tasks/{taskId}/complete", "POST /cases/{id}/documents/{documentId}", "POST /cases/{id}/participants"],
            "functions": [
                "Create evidence-rich cases with categories, priorities, and outcomes.",
                "Add tasks, complete tasks, attach documents, and invite internal or external participants.",
                "Maintain a full event timeline for non-linear case work.",
            ],
            "e2e": [
                "Create a new case from the Cases page and confirm it appears in the summary/list.",
                "Open the case, add a task, complete it, and verify the event history updates.",
                "Attach a document and confirm the attachment appears in the case detail drawer.",
                "Invite a participant, close the case with an outcome, and confirm status shifts to CLOSED.",
            ],
            "evidence": ["cases, case_tasks, case_participants, case_attachments, and case_events tables updated."]
        },
        {
            "title": "7.8 Carbon-Aware Storage And Sustainability",
            "route": "/sustainability",
            "apis": ["GET /sustainability/snapshot", "POST /sustainability/snapshot/refresh", "GET /sustainability/trend"],
            "functions": [
                "Composite Green Index visibility.",
                "Metrics for bytes stored, deduplicated, archived, kWh saved, and CO2 avoided.",
                "Daily snapshots and trend analysis for storage posture.",
            ],
            "e2e": [
                "Open Sustainability and verify the hero card, KPI cards, and trend chart render.",
                "Trigger Refresh and confirm a new snapshot timestamp is visible.",
                "Inspect trend after repeated refreshes or scheduled execution to confirm the history grows.",
            ],
            "evidence": ["sustainability_metrics table populated; calculations use dedup and archive ratios."]
        },
        {
            "title": "7.9 Federated Search",
            "route": "/federated",
            "apis": ["GET /search/federated", "Document-service federated routes under /federated/**"],
            "functions": [
                "Cross-source search against remote repositories such as email archives or legacy stores.",
                "Source badges indicating live remote, indexed, or in-platform results.",
                "Unified retrieval without forcing full migration first.",
            ],
            "e2e": [
                "Search a term known to exist in federated sources and verify mixed-result badges render correctly.",
                "Open a live remote result and confirm the user can distinguish it from an in-platform document.",
                "Review audit log for federated search/view events.",
            ],
        },
        {
            "title": "7.10 Projects, Governance, Plugins, And Industry Templates",
            "route": "/projects, /trust-center, /marketplace, /industry, /apps",
            "apis": ["GET/POST /projects/**", "GET/POST /governance/**", "GET/POST /auth/plugins/**", "GET /auth/industry-templates/**"],
            "functions": [
                "Project creation, membership, scoped defaults, and hierarchical projects.",
                "AI governance policies by global or project scope.",
                "Plugin activation and industry template discovery.",
                "Application/catalog views for installed capabilities.",
            ],
            "e2e": [
                "Create a project, add a member with scoped permissions, and verify project-specific access.",
                "Toggle AI or update a governance policy and confirm the Trust Center reflects the effective policy.",
                "Open Marketplace and activate/deactivate a plugin, then verify the project plugin list updates.",
                "Browse Industry templates and confirm templates filter by industry/category.",
            ],
        },
        {
            "title": "7.11 Workflow Engine And Designer",
            "route": "/workflow and /workflow-designer",
            "apis": ["GET/POST /workflow/definitions", "POST /workflow/instances", "POST /workflow/instances/{id}/transition", "POST /workflow/instances/{id}/approve", "POST /workflows/forwards"],
            "functions": [
                "Design workflow definitions with states, transitions, approvals, and escalations.",
                "Start workflow instances against documents.",
                "Approve, reject, cancel, peer review, and forward tasks.",
                "Track history and overdue states.",
            ],
            "e2e": [
                "Create or open a workflow definition in Designer and confirm states/transitions save.",
                "Start a workflow on a document and verify the instance is visible in Workflow.",
                "Approve the workflow as an approver or reject and resubmit as needed.",
                "Forward a workflow task to another user and confirm it appears in that user’s queue.",
            ],
            "evidence": ["workflow definitions and instances persisted; audit and notification events emitted."]
        },
        {
            "title": "7.12 Notifications And Real-Time Presence",
            "route": "/notifications",
            "apis": ["GET /notification/my", "GET /notification/my/unread", "POST /notification/{id}/read", "GET/POST/DELETE /notifications/presence/{documentId}"],
            "functions": [
                "Per-user notification inbox and unread counts.",
                "Mark-read and read-all operations.",
                "Live presence when multiple users open the same document.",
            ],
            "e2e": [
                "Trigger a workflow, agentic, or forwarding event and confirm a notification appears for the target user.",
                "Mark the notification as read and confirm unread counters decrease.",
                "Open one document in two browser sessions and confirm presence indicators appear and disappear correctly.",
            ],
        },
        {
            "title": "7.13 Email Ingestion",
            "route": "/email-ingestion",
            "apis": ["GET/POST /email-ingestion/configs", "POST /email-ingestion/configs/{id}/rules", "POST /email-ingestion/configs/{id}/poll"],
            "functions": [
                "Mailbox configuration management.",
                "Rule-based ingestion decisions.",
                "Manual polling and scheduled email harvesting into documents.",
            ],
            "e2e": [
                "Create or edit a mailbox config, add a rule, and trigger Poll Now.",
                "Verify messages matching rules become documents or intake items.",
                "Inspect document metadata to confirm the source mailbox/config is captured.",
            ],
        },
        {
            "title": "7.14 SAP Integration",
            "route": "/sap",
            "apis": ["POST /sap/invoices/process/{documentId}", "POST /sap/assets/{equipmentId}/link/{documentId}", "GET /sap-mock/purchase-orders", "GET /sap-mock/assets", "GET /sap-mock/archivelink/get"],
            "functions": [
                "Procure-to-pay invoice matching and parking.",
                "Meter-to-cash / asset-document linkage.",
                "ArchiveLink-style document retrieval via SAP object keys.",
                "Bridge-agent assisted SAP draft preparation.",
            ],
            "e2e": [
                "Use a seeded purchase order and a test invoice mentioning the same PO number and amount, then run invoice processing.",
                "Verify the result is matched and parked in the SAP mock flow or disputed when the PO is absent.",
                "Link a document to a seeded equipment ID and confirm ArchiveLink retrieval shows the document.",
            ],
            "evidence": ["sap_purchase_orders, sap_invoices, sap_assets, and sap_object_links updated."]
        },
        {
            "title": "7.15 Vendor Portals",
            "route": "/vendor-portals and /public/vendor-portal/{token}",
            "apis": ["POST /vendor-portals", "GET /vendor-portals", "Public portal routes under /public/vendor-portal/**"],
            "functions": [
                "Provision external vendor upload portals.",
                "Track required documents, compliance state, tokenised access, and watermark/audit controls.",
                "Receive external submissions into the managed corpus.",
            ],
            "e2e": [
                "Create a vendor portal from the admin view and copy the generated public URL/token.",
                "Open the public link in an incognito window, upload the required document, and submit.",
                "Return to the admin portal and confirm compliance status, upload record, and resulting document are visible.",
            ],
        },
        {
            "title": "7.16 Retention, Jurisdiction Rules, And Disposition",
            "route": "/retention and /compliance",
            "apis": ["GET/POST /retention/policies", "GET /retention/dispositions", "POST /retention/dispositions/{id}/approve", "POST /retention/scan", "GET /retention/jurisdictions/**"],
            "functions": [
                "Retention policy lifecycle.",
                "Jurisdiction/legal-framework browsing and rule checks.",
                "Disposition queue approval, rejection, and hold operations.",
            ],
            "e2e": [
                "Create or update a retention policy and assign it to a document.",
                "Trigger a retention scan and confirm eligible items appear in the disposition queue.",
                "Approve, reject, or hold a disposition and verify the new status plus audit entry.",
                "Browse jurisdiction rules in Compliance and verify category/jurisdiction filters work.",
            ],
        },
        {
            "title": "7.17 Privacy And POTRAZ Compliance",
            "route": "/compliance",
            "apis": ["GET /compliance/potraz/dashboard", "GET/POST /compliance/consents", "GET/POST /compliance/dsr", "GET/POST /compliance/breaches", "GET/POST /compliance/transfers", "GET/POST /compliance/dst", "GET /compliance/pii/inventory", "POST /compliance/snapshots"],
            "functions": [
                "Consent register and withdrawal.",
                "Data subject request workflow.",
                "Breach logging and regulator notification support.",
                "Cross-border transfer register, DST, PII inventory, and audit snapshots.",
            ],
            "e2e": [
                "Create a consent and then withdraw it.",
                "Create a DSR, move it through statuses, and execute it.",
                "Log a breach and trigger POTRAZ notification.",
                "Generate a compliance snapshot and confirm it is retrievable from the snapshots list.",
            ],
        },
        {
            "title": "7.18 Audit",
            "route": "/audit",
            "apis": ["GET /audit/recent", "GET /audit/stats", "GET /audit/logs/by-user/**", "GET /audit/logs/by-action/**", "GET /audit/logs/by-resource/**"],
            "functions": [
                "Immutable event history across document, workflow, AI, and admin operations.",
                "Filtering by user, action, resource, resource type, and date range.",
                "High-level audit statistics.",
            ],
            "e2e": [
                "Perform a known action such as upload, approve, or accept an agent suggestion.",
                "Open Audit and filter by the corresponding action or resource.",
                "Confirm the row includes actor, timestamp, action, and details.",
            ],
        },
        {
            "title": "7.19 Analytics And Predictions",
            "route": "/analytics and /predictions",
            "apis": ["Analytics are populated from cross-service operational data."],
            "functions": [
                "Operational charts and workflow-health views.",
                "Predictive and anomaly-oriented dashboards for privileged roles.",
            ],
            "e2e": [
                "Open Analytics and verify charts render using live data.",
                "Open Workflow Health and confirm overdue or bottleneck views update after creating test workflows.",
                "Open Predictions as an allowed role and confirm model outputs render without authorization errors.",
            ],
        },
        {
            "title": "7.20 Knowledge Graph",
            "route": "/knowledge-graph",
            "apis": ["Knowledge graph view is built from document and relationship data."],
            "functions": [
                "Graph visualisation of document, project, and entity relationships.",
                "Exploration of linked content and potentially inferred associations.",
            ],
            "e2e": [
                "Create or verify linked documents, projects, or cases.",
                "Open Knowledge Graph and confirm nodes and edges reflect those relationships.",
                "Select a node and verify drill-through into the underlying record.",
            ],
        },
        {
            "title": "7.21 Blockchain Notarization",
            "route": "/documents (document detail actions)",
            "apis": ["POST /documents/{id}/notarize", "GET /documents/{id}/verify"],
            "functions": [
                "Immutable document fingerprinting using chained block hashes.",
                "Verification that current content still matches the last notarised state.",
            ],
            "e2e": [
                "Notarize a document and verify a positive chain result.",
                "Change the document content or hash and rerun verification to confirm tamper detection.",
            ],
            "evidence": ["document_notarizations records show block number, content hash, previous hash, and block hash."]
        },
    ]

    for module in modules:
        add_module_section(doc, module["title"], module["route"], module.get("apis", []), module["functions"], module["e2e"], module.get("evidence"))


def add_appendix(doc: Document):
    doc.add_heading("8. Reusable API Smoke Commands", level=1)
    code_block(doc, [
        f"curl -H \"Authorization: Bearer $TOKEN\" {GATEWAY_URL}/documents/my | jq '.data[0]'",
        f"curl -H \"Authorization: Bearer $TOKEN\" {GATEWAY_URL}/agentic/team | jq",
        f"curl -H \"Authorization: Bearer $TOKEN\" {GATEWAY_URL}/cases/summary | jq",
        f"curl -H \"Authorization: Bearer $TOKEN\" {GATEWAY_URL}/sustainability/snapshot | jq",
        f"curl -H \"Authorization: Bearer $TOKEN\" {GATEWAY_URL}/audit/stats | jq",
    ])

    doc.add_heading("9. Database Verification Shortcuts", level=1)
    code_block(doc, [
        f"{POSTGRES_CMD} -c \"SELECT id,title,classification_label,language FROM documents ORDER BY created_at DESC LIMIT 5;\"",
        f"{POSTGRES_CMD} -c \"SELECT agent_name,intent_type,status,created_at FROM agentic_suggestions ORDER BY created_at DESC LIMIT 10;\"",
        f"{POSTGRES_CMD} -c \"SELECT id,title,status,priority,created_at FROM cases ORDER BY created_at DESC LIMIT 10;\"",
        f"{POSTGRES_CMD} -c \"SELECT captured_at,green_index,kwh_saved,co2_kg_avoided FROM sustainability_metrics ORDER BY captured_at DESC LIMIT 10;\"",
        f"{POSTGRES_CMD} -c \"SELECT action,username,details FROM audit_log ORDER BY created_at DESC LIMIT 20;\"",
    ])

    doc.add_heading("10. Recommended End-To-End Regression Pack", level=1)
    for item in [
        "Authentication: login, logout, role change, access denial.",
        "Document lifecycle: upload, preview, search, workflow, approval, share, audit.",
        "AI lifecycle: classify, GDPR scan, agent suggestion, accept/dismiss, audit.",
        "Operational extensions: SAP process, vendor portal submission, email poll, federated query.",
        "2026 features: multi-agent runs, case creation/closure, sustainability refresh, Shona/Ndebele language detection.",
    ]:
        bullet(doc, item)


def main():
    doc = Document()
    style_document(doc)
    add_title_page(doc)
    add_intro(doc)
    add_module_catalog(doc)
    build_modules(doc)
    add_appendix(doc)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
